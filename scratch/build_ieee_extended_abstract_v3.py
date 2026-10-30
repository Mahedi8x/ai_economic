# -*- coding: utf-8 -*-
"""
Generate Updated, Publication-Grade IEEE CS BDC Symposium 2026 Extended Abstract (v3).
Title updated per user request:
"Explainable Machine Learning and a Grounded Policy Retrieval Prototype for Food Inflation Forecasting and Economic Early Warning in Bangladesh"

Features:
- Exactly 2 pages in PDF rendering (A4 format, Book Antiqua font).
- Authors: Md Mohi Uddin Mahedi (Primary, 1st Author), Mst. Mimuna Parvin Mim (Co-Author, 2nd Author).
- Generates both Author and Anonymized versions (DOCX & PDF).
- Copies to Desktop and repo base directory.
- Reconciled numbers:
  * Conformal coverage: 83.6% (rice), 76.0% (lentils), 75.2% (flour) over 1,368 days.
  * Table I daily RF: "—" (not evaluated / not applicable).
  * Dataset: 6,117 observations (Jan 2010 - Sep 2026), 201 months. Test n ≈ 33.
"""

import os
import subprocess
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(r"C:\ai_economic")
DESKTOP_DIR = Path(r"C:\Users\Mehedi\Desktop")
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

TITLE_TEXT = "Explainable Machine Learning and a Grounded Policy Retrieval Prototype for Food Inflation Forecasting and Economic Early Warning in Bangladesh"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def build_docx(anonymized=False):
    doc = Document()
    
    # Page setup: A4, Top/Bottom 1.0", Left/Right 0.75"
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Book Antiqua'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)
    normal_style.paragraph_format.line_spacing = 1.05
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(3)
    run_title = p_title.add_run(TITLE_TEXT)
    run_title.font.name = 'Book Antiqua'
    run_title.font.size = Pt(12.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle / Conference Header
    p_conf = doc.add_paragraph()
    p_conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_conf.paragraph_format.space_after = Pt(5)
    run_conf = p_conf.add_run("EXTENDED ABSTRACT — IEEE CS BDC SYMPOSIUM 2026")
    run_conf.font.name = 'Book Antiqua'
    run_conf.font.size = Pt(9.0)
    run_conf.font.italic = True
    run_conf.font.color.rgb = RGBColor(71, 85, 105)

    # Author Block
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(6)
    if anonymized:
        run_auth = p_author.add_run("Anonymous Author(s) for Double-Blind Peer Review\n(Author names, institutional affiliations, and repository links withheld in compliance with double-blind review policy)")
        run_auth.font.italic = True
        run_auth.font.size = Pt(8.5)
    else:
        run_auth = p_author.add_run("Md Mohi Uddin Mahedi¹*, Mst. Mimuna Parvin Mim²\nDepartment of Computer Science and Engineering\nOpen-Source Codebase: https://github.com/Mahedi8x/ai_economic")
        run_auth.font.bold = True
        run_auth.font.size = Pt(8.8)

    # Keywords
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_after = Pt(5)
    r_kwh = p_kw.add_run("Keywords— ")
    r_kwh.font.bold = True
    r_kwh.font.size = Pt(9.0)
    r_kwt = p_kw.add_run("Food inflation forecasting, Explainable machine learning, Economic early warning, Conformal prediction, Policy retrieval.")
    r_kwt.font.italic = True
    r_kwt.font.size = Pt(9.0)

    # Headings
    def add_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.font.name = 'Book Antiqua'
        r.font.size = Pt(10.0)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return p

    # 1. Research Area
    add_heading("1. Research Area")
    doc.add_paragraph(
        "Applied machine learning and computational macroeconomics (time-series forecasting). Research question: can "
        "gradient-boosted models trained on publication-lag-aligned daily market data and monthly macro indicators forecast "
        "Bangladesh's food inflation better than linear and naive-persistence baselines, and are the resulting early-warning "
        "signals reliable enough to be useful to policymakers?"
    )

    # 2. Goals/Objectives
    add_heading("2. Goals/Objectives")
    doc.add_paragraph(
        "Official food CPI and reserve figures from the Bangladesh Bureau of Statistics (BBS) and Bangladesh Bank appear "
        "around Day 7–10 of the following month. This delay is modest, but during shock regimes (currency depreciation, "
        "commodity spikes) it leaves monetary and procurement authorities without a timely signal. We therefore test, as an "
        "honest benchmark rather than a claim of superiority: (i) whether gradient boosting improves one-month and three-month-ahead "
        "food inflation forecasts over linear regression and naive persistence; (ii) how well rolling conformal prediction "
        "calibrates uncertainty bands for retail staple prices; (iii) whether unsupervised anomaly detection gives advance "
        "warning of macroeconomic stress episodes without excessive false alarms; and (iv) how a policy retrieval pipeline "
        "prototype performs on grounding domain economic inquiries on central bank publications."
    )

    # 3. Methodology
    add_heading("3. Methodology")
    doc.add_paragraph(
        "Data. A daily panel (January 2010–September 2026, 6,117 observations, 41 indicators) is combined with a monthly "
        "panel of official Bangladesh Bank/BBS series (201 months, no interpolation). Release schedules (BBS CPI on Day 10 "
        "of M+1, Bangladesh Bank reserves on Day 7 of M+1, WFP retail surveys on Day 15 of M) are enforced by step-function "
        "forward fills; daily international series (Brent crude, CBOT wheat, US Dollar Index, natural gas, gold) enter "
        "contemporaneously. Known limitation: annual World Bank indicators are step-filled by calendar year and are not lagged "
        "to their release dates, which introduces look-ahead bias that we have not yet removed. Forecasting. XGBoost is trained "
        "on chronological splits (train ≤ 2021, validation 2022–2023, test 2024–2026) and compared with Ridge regression, "
        "Random Forest, and naive persistence using Harvey-Leybourne-Newbold corrected Diebold-Mariano (DM) tests; feature attributions "
        "use TreeSHAP. Uncertainty. For coarse rice, lentils, and wheat flour, rolling walk-forward conformal prediction "
        "(730-day calibration window) targets 80% intervals; time-series dependence means no formal coverage guarantee holds. "
        "Early warning. An Isolation Forest is fitted only on pre-2022 data, combined with a resetting CUSUM filter, and evaluated "
        "out-of-sample (2022–2026) against documented stress episodes. Retrieval. A ChromaDB dense-retrieval prototype indexes "
        "authentic policy summaries (Bangladesh Bank, BBS, IMF, TCB)."
    )

    # 4. Results and Analysis
    add_heading("4. Results and Analysis")
    doc.add_paragraph(
        "Forecasting. On the monthly test set (2024–2026, n ≈ 33), XGBoost lowers RMSE relative to Ridge by 20.0% at t+1 and "
        "32.5% at t+3, but neither difference is significant at the 5% level (DM p = 0.068, 0.058). XGBoost does not beat naive "
        "persistence (t+1: 1.171% vs. 1.160%) or Random Forest (Table I). In daily 30-day-ahead forecasting it clearly beats "
        "Ridge (DM = 3.74, p < 0.001) yet stays slightly behind persistence (1.158% vs. 1.137%); because the target is a step-filled "
        "monthly series, the effective sample is small and this p-value is optimistic. TreeSHAP ranks reserve adequacy "
        "(months of imports, 19.3%), remittances (15.9%), and headline-inflation persistence (15.4%) highest, whereas Brent crude "
        "(1.4%) and wheat (1.2%) contribute little; since the top feature is an annual series, this ranking is itself affected by "
        "the look-ahead issue above."
    )

    # Table I
    tbl = doc.add_table(rows=4, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Horizon", "Naive", "Ridge", "Random Forest", "XGBoost (vs. Ridge)", "DM vs. Ridge (p)"]
    hdr_row = tbl.rows[0]
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.text = text
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=45, bottom=45, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Book Antiqua'
            run.font.size = Pt(8.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ["t+1 (1 month)", "1.160%", "1.464%", "1.087%", "1.171% (-20.0%)", "1.82 (0.068)"],
        ["t+3 (3 months)", "1.757%", "2.578%", "1.725%", "1.740% (-32.5%)", "1.90 (0.058)"],
        ["30-day daily", "1.137%", "2.457%", "—", "1.158% (-52.9%)", "3.74 (<0.001)"]
    ]

    for r_idx, row_vals in enumerate(data_rows):
        row = tbl.rows[r_idx + 1]
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=35, bottom=35, left=50, right=50)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Book Antiqua'
                run.font.size = Pt(8.0)
                if c_idx == 4 or c_idx == 5:
                    run.font.bold = True

    p_tbl_lbl = doc.add_paragraph()
    p_tbl_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_lbl.paragraph_format.space_before = Pt(2)
    p_tbl_lbl.paragraph_format.space_after = Pt(3)
    r_tl = p_tbl_lbl.add_run("Table I: Out-of-sample RMSE and Diebold-Mariano tests (lowest RMSE per row underlined).")
    r_tl.font.size = Pt(7.5)
    r_tl.font.italic = True
    r_tl.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph(
        "Uncertainty and early warning. Over 1,368 out-of-sample days, conformal intervals reached 83.6% empirical coverage "
        "for coarse rice but only 76.0% (lentils) and 75.2% (flour) against the 80% target. The anomaly detector flagged both "
        "out-of-sample episodes (2022 FX drain, 2024 political shock) 80 and 90 days before onset, but it also flagged 74% of "
        "out-of-sample days. Most of these days lie outside the pre-2022 training range (e.g., USD/BDT), and the lead time is bounded "
        "by a 90-day search window, so it should be read as a distribution-shift signal. Precision is 56.5% against a 42% crisis-day "
        "base rate (F1 0.72 vs. 0.59 for always alarming), so alert-persistence filtering is required before any operational use. "
        "Retrieval. The prototype reached 82.0% mean keyword coverage on 10 queries over four authentic policy summaries; LLM "
        "generation and hallucination have not yet been evaluated, so no grounding claim is made. Open issues. Calendar-year look-ahead "
        "in annual indicators; a short test period (≈33 months) with two crisis episodes; hand-labelled crisis windows; incomplete "
        "conformal calibration for lentils and flour; and absence of a generation-level retrieval evaluation. Code and data are released openly."
    )

    # 5. References
    add_heading("5. References")
    refs = [
        "[1] T. Chen and C. Guestrin, \"XGBoost: A scalable tree boosting system,\" in Proc. 22nd ACM SIGKDD Int. Conf. Knowl. Discovery Data Mining, 2016, pp. 785–794.",
        "[2] F. T. Liu, K. M. Ting, and Z.-H. Zhou, \"Isolation forest,\" in Proc. 8th IEEE Int. Conf. Data Mining (ICDM), 2008, pp. 413–422.",
        "[3] F. X. Diebold and R. S. Mariano, \"Comparing predictive accuracy,\" J. Business & Economic Statistics, vol. 13, no. 3, pp. 253–263, 1995."
    ]
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(1.5)
        r_ref = p_ref.add_run(ref)
        r_ref.font.name = 'Book Antiqua'
        r_ref.font.size = Pt(8.0)

    return doc

def generate_abstract_html(anonymized=False):
    author_block = """
        <div class="author-block">
            <strong>Anonymous Author(s) for Double-Blind Peer Review</strong><br>
            <span style="font-size: 8pt; color: #64748b; font-style: italic;">
                (Author names, institutional affiliations, and repository links withheld in compliance with double-blind review policy)
            </span>
        </div>
    """ if anonymized else """
        <div class="author-block">
            <strong>Md Mohi Uddin Mahedi<sup>1*</sup>, Mst. Mimuna Parvin Mim<sup>2</sup></strong><br>
            <span style="font-size: 8.2pt; color: #334155;">
                Department of Computer Science and Engineering<br>
                Open-Source Repository: <u>https://github.com/Mahedi8x/ai_economic</u>
            </span>
        </div>
    """

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Extended Abstract - IEEE CS BDC Symposium 2026</title>
    <style>
        @page {{
            size: A4;
            margin: 24mm 19mm 24mm 19mm; /* 0.95" top/bottom, 0.75" left/right */
        }}

        body {{
            font-family: 'Book Antiqua', Palatino, 'Palatino Linotype', Georgia, serif;
            color: #0f172a;
            line-height: 1.32;
            font-size: 9.15pt;
            margin: 0;
            padding: 0;
            background: #ffffff;
            text-align: justify;
        }}

        .header {{
            text-align: center;
            margin-bottom: 7px;
        }}

        h1.title {{
            font-size: 12.2pt;
            font-weight: bold;
            color: #0f172a;
            margin: 0 0 3px 0;
            line-height: 1.22;
            text-align: center;
        }}

        .conf-badge {{
            font-size: 8.3pt;
            font-style: italic;
            color: #475569;
            margin-bottom: 3px;
            text-transform: uppercase;
            letter-spacing: 0.3px;
        }}

        .author-block {{
            text-align: center;
            margin: 3px 0 5px 0;
            font-size: 8.5pt;
        }}

        .keywords {{
            margin-bottom: 7px;
            font-size: 8.8pt;
        }}

        .keywords strong {{
            font-weight: bold;
            color: #1e293b;
        }}

        h2.sec-heading {{
            font-size: 9.6pt;
            font-weight: bold;
            color: #1e3a8a;
            margin-top: 6px;
            margin-bottom: 2px;
            page-break-after: avoid;
        }}

        p {{
            margin: 0 0 5px 0;
            text-indent: 0;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 5px 0 2px 0;
            font-size: 7.8pt;
            page-break-inside: avoid;
        }}

        th {{
            background-color: #1e3a8a;
            color: #ffffff;
            font-weight: bold;
            padding: 3.5px 5px;
            border: 1px solid #94a3b8;
            text-align: center;
        }}

        td {{
            padding: 3px 5px;
            border: 1px solid #cbd5e1;
            text-align: center;
        }}

        tr:nth-child(even) {{
            background-color: #f8fafc;
        }}

        .table-caption {{
            font-size: 7.2pt;
            font-style: italic;
            color: #475569;
            text-align: center;
            margin-bottom: 5px;
        }}

        .ref-item {{
            margin-bottom: 2px;
            font-size: 7.6pt;
            line-height: 1.25;
        }}
    </style>
</head>
<body>

    <div class="header">
        <h1 class="title">{TITLE_TEXT}</h1>
        <div class="conf-badge">EXTENDED ABSTRACT — IEEE CS BDC SYMPOSIUM 2026</div>
        {author_block}
    </div>

    <div class="keywords">
        <strong>Keywords—</strong> <em>Food inflation forecasting, Explainable machine learning, Economic early warning, Conformal prediction, Policy retrieval.</em>
    </div>

    <h2 class="sec-heading">1. Research Area</h2>
    <p>
        Applied machine learning and computational macroeconomics (time-series forecasting). Research question: can gradient-boosted models trained on publication-lag-aligned daily market data and monthly macro indicators forecast Bangladesh's food inflation better than linear and naive-persistence baselines, and are the resulting early-warning signals reliable enough to be useful to policymakers?
    </p>

    <h2 class="sec-heading">2. Goals/Objectives</h2>
    <p>
        Official food CPI and reserve figures from the Bangladesh Bureau of Statistics (BBS) and Bangladesh Bank appear around Day 7–10 of the following month. This delay is modest, but during shock regimes (currency depreciation, commodity spikes) it leaves monetary and procurement authorities without a timely signal. We therefore test, as an honest benchmark rather than a claim of superiority: (i) whether gradient boosting improves one-month and three-month-ahead food inflation forecasts over linear regression and naive persistence; (ii) how well rolling conformal prediction calibrates uncertainty bands for retail staple prices; (iii) whether unsupervised anomaly detection gives advance warning of macroeconomic stress episodes without excessive false alarms; and (iv) how a policy retrieval pipeline prototype performs on domain queries.
    </p>

    <h2 class="sec-heading">3. Methodology</h2>
    <p>
        <strong>Data.</strong> A daily panel (January 2010–September 2026, 6,117 observations, 41 indicators) is combined with a monthly panel of official Bangladesh Bank/BBS series (201 months, no interpolation). Release schedules (BBS CPI on Day 10 of M+1, Bangladesh Bank reserves on Day 7 of M+1, WFP retail surveys on Day 15 of M) are enforced by step-function forward fills; daily international series (Brent crude, CBOT wheat, US Dollar Index, natural gas, gold) enter contemporaneously. Known limitation: annual World Bank indicators are step-filled by calendar year and are not lagged to their release dates, which introduces look-ahead bias that we have not yet removed. <strong>Forecasting.</strong> XGBoost [1] is trained on chronological splits (train ≤ 2021, validation 2022–2023, test 2024–2026) and compared with Ridge regression, Random Forest and naive persistence using Harvey-Leybourne-Newbold corrected Diebold-Mariano (DM) tests [3]; feature attributions use TreeSHAP. <strong>Uncertainty.</strong> For coarse rice, lentils and wheat flour, rolling walk-forward conformal prediction (730-day calibration window) targets 80% intervals; time-series dependence means no formal coverage guarantee holds. <strong>Early warning.</strong> An Isolation Forest [2] is fitted only on pre-2022 data, combined with a resetting CUSUM filter, and evaluated out-of-sample (2022–2026) against documented stress episodes. <strong>Retrieval.</strong> A ChromaDB dense-retrieval prototype indexes four policy summaries (Bangladesh Bank, BBS, IMF, TCB).
    </p>

    <h2 class="sec-heading">4. Results and Analysis</h2>
    <p>
        <strong>Forecasting.</strong> On the monthly test set (2024–2026, n ≈ 33), XGBoost lowers RMSE relative to Ridge by 20.0% at t+1 and 32.5% at t+3, but neither difference is significant at the 5% level (DM p = 0.068, 0.058). XGBoost does not beat naive persistence (t+1: 1.171% vs. 1.160%) or Random Forest (Table I). In daily 30-day-ahead forecasting it clearly beats Ridge (DM = 3.74, p &lt; 0.001) yet stays slightly behind persistence (1.158% vs. 1.137%); because the target is a step-filled monthly series, the effective sample is small and this p-value is optimistic. TreeSHAP ranks reserve adequacy (months of imports, 19.3%), remittances (15.9%) and headline-inflation persistence (15.4%) highest, whereas Brent crude (1.4%) and wheat (1.2%) contribute little; since the top feature is an annual series, this ranking is itself affected by the look-ahead issue above.
    </p>

    <table>
        <thead>
            <tr>
                <th>Horizon</th>
                <th>Naive</th>
                <th>Ridge</th>
                <th>Random Forest</th>
                <th>XGBoost (vs. Ridge)</th>
                <th>DM vs. Ridge (p)</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td>t+1 (1 month)</td>
                <td>1.160%</td>
                <td>1.464%</td>
                <td><u>1.087%</u></td>
                <td>1.171% (-20.0%)</td>
                <td>1.82 (0.068)</td>
            </tr>
            <tr>
                <td>t+3 (3 months)</td>
                <td>1.757%</td>
                <td>2.578%</td>
                <td><u>1.725%</u></td>
                <td>1.740% (-32.5%)</td>
                <td>1.90 (0.058)</td>
            </tr>
            <tr>
                <td>30-day daily</td>
                <td><u>1.137%</u></td>
                <td>2.457%</td>
                <td>—</td>
                <td><strong>1.158% (-52.9%)</strong></td>
                <td><strong>3.74 (&lt;0.001)</strong></td>
            </tr>
        </tbody>
    </table>
    <div class="table-caption">
        Table I: Out-of-sample RMSE and Diebold-Mariano tests (lowest RMSE per row underlined).
    </div>

    <p>
        <strong>Uncertainty and early warning.</strong> Over 1,368 out-of-sample days, conformal intervals reached 83.6% empirical coverage for coarse rice but only 76.0% (lentils) and 75.2% (flour) against the 80% target. The anomaly detector flagged both out-of-sample episodes (2022 FX drain, 2024 political shock) 80 and 90 days before onset, but it also flagged 74% of out-of-sample days. Most of these days lie outside the pre-2022 training range (e.g., USD/BDT), and the lead time is bounded by a 90-day search window, so it should be read as a distribution-shift signal. Precision is 56.5% against a 42% crisis-day base rate (F1 0.72 vs. 0.59 for always alarming), so alert-persistence filtering is required before any operational use. <strong>Retrieval.</strong> The prototype reached 82.0% mean keyword coverage on 10 queries over four policy summaries; LLM generation and hallucination have not yet been evaluated, so no grounding claim is made.
    </p>

    <p>
        <strong>Open issues.</strong> Calendar-year look-ahead in annual indicators; a short test period (≈33 months) with two crisis episodes; hand-labelled crisis windows; incomplete conformal calibration for lentils and flour; and absence of a generation-level retrieval evaluation. Code and data are released openly; corrected results will be versioned in the repository.
    </p>

    <h2 class="sec-heading">5. References</h2>
    <div class="ref-item">
        [1] T. Chen and C. Guestrin, "XGBoost: A scalable tree boosting system," in <em>Proc. 22nd ACM SIGKDD Int. Conf. Knowl. Discovery Data Mining</em>, 2016, pp. 785–794.
    </div>
    <div class="ref-item">
        [2] F. T. Liu, K. M. Ting, and Z.-H. Zhou, "Isolation forest," in <em>Proc. 8th IEEE Int. Conf. Data Mining (ICDM)</em>, 2008, pp. 413–422.
    </div>
    <div class="ref-item">
        [3] F. X. Diebold and R. S. Mariano, "Comparing predictive accuracy," <em>J. Business & Economic Statistics</em>, vol. 13, no. 3, pp. 253–263, 1995.
    </div>

</body>
</html>
"""

def main():
    print("[1/4] Generating DOCX files...")
    doc_anon = build_docx(anonymized=True)
    doc_auth = build_docx(anonymized=False)
    
    anon_docx = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract_Anonymized.docx"
    auth_docx = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract.docx"
    
    doc_anon.save(anon_docx)
    doc_auth.save(auth_docx)
    print(f" [+] Saved: {anon_docx}")
    print(f" [+] Saved: {auth_docx}")

    print("[2/4] Rendering publication-grade PDFs via Edge Headless...")
    # Anonymized HTML & PDF
    anon_html = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract_Anonymized.html"
    anon_pdf = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract_Anonymized.pdf"
    with open(anon_html, "w", encoding="utf-8") as f:
        f.write(generate_abstract_html(anonymized=True))
    
    subprocess.run([
        EDGE_PATH, "--headless", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={anon_pdf}", str(anon_html)
    ], capture_output=True, text=True)

    # Author HTML & PDF
    auth_html = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract.html"
    auth_pdf = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract.pdf"
    with open(auth_html, "w", encoding="utf-8") as f:
        f.write(generate_abstract_html(anonymized=False))

    subprocess.run([
        EDGE_PATH, "--headless", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={auth_pdf}", str(auth_html)
    ], capture_output=True, text=True)

    print(f" [+] Rendered: {anon_pdf}")
    print(f" [+] Rendered: {auth_pdf}")

    print("[3/4] Copying all files to Desktop...")
    for f in [anon_docx, auth_docx, anon_pdf, auth_pdf]:
        shutil.copy2(f, DESKTOP_DIR / f.name)
        print(f" [+] Copied to Desktop: {DESKTOP_DIR / f.name}")

    print("[4/4] Verifying exact page counts...")
    import pypdf
    for p in [anon_pdf, auth_pdf]:
        reader = pypdf.PdfReader(str(p))
        num_pages = len(reader.pages)
        print(f" [+] {p.name} -> Exact Page Count: {num_pages}")
        if num_pages != 2:
            print(f" [!] WARNING: {p.name} is {num_pages} pages, expected 2!")

if __name__ == "__main__":
    main()
