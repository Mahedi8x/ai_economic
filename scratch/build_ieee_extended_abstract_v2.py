# -*- coding: utf-8 -*-
"""
Generate Corrected, Honest, Publication-Grade IEEE CS BDC Symposium 2026 Extended Abstract.
Incorporate all 9 reviewer corrections:
1. SHAP truth: reserves (19.3%), remittance (15.9%), inflation persistence (15.4%), oil/wheat minor.
2. Honest performance: XGBoost beats Ridge by 20.0-32.5%, but not significant at 5% (p=0.068, 0.058).
   RF and Naive are competitive. Daily beats Ridge (p<0.001), on par with Naive.
3. BBS delay corrected to ~1-2 weeks (Day 10 of M+1).
4. Anomaly detection: 80 and 90 days lead time, but acknowledges 74% alarm rate needing CUSUM.
5. Conformal: Rice achieves 85.1% OOS coverage, lentils (74.8%) and flour (75.1%) below 80% target.
6. RAG: Audit on keyword/semantic coverage (82% on 10 queries); LLM/hallucination not yet evaluated.
7. Data panel: Jan 2010 to September 2026 (201 months), annual feature look-ahead concern acknowledged.
8. Daily RF omitted/noted; numbers strictly match CSVs.
9. Dashboard described as "prototype dashboard".
Authors:
1. Md Mohi Uddin Mahedi (Primary)
2. Mst. Mimuna Parvin Mim (Co-Author)
"""

import os
import subprocess
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(r"C:\ai_economic")
DESKTOP_DIR = Path(r"C:\Users\Mehedi\Desktop")
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def build_docx(anonymized=False):
    doc = Document()
    
    # Page setup: A4, Top/Bottom 1.0", Left/Right 0.75"
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Book Antiqua'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)
    normal_style.paragraph_format.line_spacing = 1.05
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(3)
    run_title = p_title.add_run("Explainable Machine Learning and Grounded Policy Retrieval Framework for Food Inflation Forecasting and Economic Early Warning in Bangladesh")
    run_title.font.name = 'Book Antiqua'
    run_title.font.size = Pt(13.0)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle / Conference Header
    p_conf = doc.add_paragraph()
    p_conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_conf.paragraph_format.space_after = Pt(5)
    run_conf = p_conf.add_run("Extended Abstract — IEEE CS BDC Symposium 2026")
    run_conf.font.name = 'Book Antiqua'
    run_conf.font.size = Pt(9.0)
    run_conf.font.italic = True
    run_conf.font.color.rgb = RGBColor(71, 85, 105)

    # Author Block
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(6)
    if anonymized:
        run_auth = p_author.add_run("Anonymous Author(s) for Double-Blind Peer Review\n(Author details withheld in compliance with double-blind review policy)")
        run_auth.font.italic = True
        run_auth.font.size = Pt(9.0)
    else:
        run_auth = p_author.add_run("Md Mohi Uddin Mahedi¹, Mst. Mimuna Parvin Mim²\nDepartment of Computer Science & Engineering\nOpen-Source Codebase: https://github.com/Mahedi8x/ai_economic")
        run_auth.font.bold = True
        run_auth.font.size = Pt(9.0)

    # Keywords
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_after = Pt(5)
    r_kwh = p_kw.add_run("Keywords— ")
    r_kwh.font.bold = True
    r_kwh.font.size = Pt(9.0)
    r_kwt = p_kw.add_run("Food Inflation Nowcasting, Explainable Machine Learning, Economic Early Warning, Conformal Prediction, Policy Information Retrieval.")
    r_kwt.font.italic = True
    r_kwt.font.size = Pt(9.0)

    # Headings
    def add_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.font.name = 'Book Antiqua'
        r.font.size = Pt(10.0)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return p

    # 1. Research Area
    add_heading("1. Research Area")
    doc.add_paragraph(
        "Applied Machine Learning, Computational Macroeconomics, Empirical Time-Series Nowcasting, and Information Retrieval."
    )

    # 2. Goals/Objectives
    add_heading("2. Goals/Objectives")
    doc.add_paragraph(
        "This research investigates macroeconomic volatility and staple food price forecasting in developing, import-dependent economies, focusing empirically on Bangladesh. In official surveillance, Consumer Price Index (CPI) and foreign exchange reserve releases from the Bangladesh Bureau of Statistics (BBS) and Bangladesh Bank typically appear around Day 7 to Day 10 of following month (Month M+1). While this represents a modest administrative delay of one to two weeks, rapid international commodity and currency fluctuations during shock regimes create an informational gap for monetary and procurement authorities. The core objective is to evaluate whether non-linear machine learning architectures can extract actionable early signals from high-frequency international markets and domestic macro indicators ahead of official monthly bulletins. Specifically, the study investigates: (i) whether gradient boosting improves short-horizon food inflation nowcasting relative to linear regression and naive persistence baselines; (ii) how empirical conformal prediction behaves when calibrating retail food price uncertainty bands; (iii) whether unsupervised anomaly detection can provide advance warning before systemic macroeconomic distress episodes; and (iv) how policy retrieval pipelines perform when grounding domain economic inquiries on central bank publications."
    )

    # 3. Methodology
    add_heading("3. Methodology")
    doc.add_paragraph(
        "The empirical architecture is anchored on a continuous daily macroeconomic panel spanning January 2010 to September 2026 (6,117 daily observations, 41 indicators) combined with an official monthly macroeconomic panel (201 continuous months, zero synthetic interpolation). Institutional release schedules (BBS CPI on Day 10 of M+1, Bangladesh Bank reserves on Day 7 of M+1, WFP retail food price surveys on Day 15 of M) are strictly enforced via step-function forward fills to prevent look-ahead bias, while daily international shocks (Brent crude oil, CBOT wheat futures, US Dollar Index, natural gas, gold) enter contemporaneously. Annual structural indicators from the World Bank (external debt, import coverage ratios) are incorporated on an as-of release basis, though annual step-filling introduces an open look-ahead timing consideration. For food inflation nowcasting, an Extreme Gradient Boosting (XGBoost) model is trained using chronological walk-forward splits and compared against linear Ridge regression, Random Forest, and naive persistence baselines using Harvey-Leybourne-Newbold corrected Diebold-Mariano tests. Feature attributions are extracted via TreeSHAP. For staple retail commodities (coarse rice, lentils, wheat flour), a rolling walk-forward adaptive conformal prediction engine is evaluated over a 730-day calibration window to generate distribution-free 80% nominal coverage intervals. For crisis surveillance, an Isolation Forest is trained exclusively on pre-2022 historical controls and evaluated out-of-sample on 2022–2026 stress regimes with a reset-capable Cumulative Sum (CUSUM) filter. A prototype dense retrieval module (ChromaDB) indexes authentic central bank policy texts."
    )

    # 4. Results and Analysis
    add_heading("4. Results and Analysis")
    doc.add_paragraph(
        "In monthly out-of-sample evaluation (2024–2026, n ≈ 33 months), XGBoost reduces root mean square error (RMSE) relative to a linear Ridge baseline by 20.0% at the t+1 horizon (1.171% vs. 1.464%) and by 32.5% at the t+3 horizon (1.740% vs. 2.578%), but neither difference is statistically significant at the 5% level (HLN-corrected Diebold-Mariano p = 0.068, 0.058), and XGBoost does not outperform naive persistence (1.160% at t+1) or Random Forest (1.087% at t+1, 1.725% at t+3) as detailed in Table I. In high-frequency daily 30-day trajectory forecasting, XGBoost substantially outperforms Ridge (RMSE 1.158% vs. 2.457%, a 52.9% reduction; DM = 3.74, p < 0.001), yet remains marginally behind naive persistence (1.137%), indicating dominant short-term price stickiness in domestic staple markets. TreeSHAP feature attribution assigns the highest importance to reserve-adequacy indicators (months of imports: 19.3%), remittance inflows (15.9%), and headline inflation persistence (15.4%), while global crude oil (1.4%) and CBOT wheat (1.2%) exhibit minor direct attribution; annual structural indicators among them remain an open look-ahead concern."
    )

    # Table I
    tbl = doc.add_table(rows=4, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Forecast Horizon", "Naive Baseline", "Linear Ridge", "Random Forest", "Proposed XGBoost", "DM Stat vs Ridge (p-val)"]
    hdr_row = tbl.rows[0]
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.text = text
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Book Antiqua'
            run.font.size = Pt(8.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ["t+1 (1 Month Ahead)", "1.160%", "1.464%", "1.087%", "1.171% (-20.0%)", "1.82 (p = 0.068)"],
        ["t+3 (3 Months Ahead)", "1.757%", "2.578%", "1.725%", "1.740% (-32.5%)", "1.90 (p = 0.058)"],
        ["30-Day Forward Daily", "1.137%", "2.457%", "—", "1.158% (-52.9%)", "3.74 (p < 0.001)"]
    ]

    for r_idx, row_vals in enumerate(data_rows):
        row = tbl.rows[r_idx + 1]
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Book Antiqua'
                run.font.size = Pt(8.0)
                if c_idx == 4 or c_idx == 5:
                    run.font.bold = True

    p_tbl_lbl = doc.add_paragraph()
    p_tbl_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_lbl.paragraph_format.space_before = Pt(2)
    p_tbl_lbl.paragraph_format.space_after = Pt(3)
    r_tl = p_tbl_lbl.add_run("Table I: Out-of-Sample Forecasting Performance & Diebold-Mariano Tests across Horizons")
    r_tl.font.size = Pt(7.5)
    r_tl.font.italic = True
    r_tl.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph(
        "In retail price uncertainty estimation, rolling walk-forward conformal prediction achieved 85.1% empirical out-of-sample coverage for coarse rice over 1,366 unseen test days, whereas lentils (74.8%) and flour (75.1%) fell short of the 80% nominal target, highlighting commodity-specific calibration disparities. The unsupervised anomaly detector flagged both out-of-sample economic crisis episodes 80 days (2022 foreign exchange drain) and 90 days (2024 political shock) ahead of onset, but flagged 74.0% of out-of-sample days due to persistent post-2022 macroeconomic disequilibrium, indicating that raw alert frequencies must be coupled with secondary persistence filters before operational deployment. The retrieval module was audited for keyword and semantic coverage, achieving an 82.0% mean grounding score across 10 central bank queries; full LLM generation and hallucination benchmarks have not yet been evaluated. A prototype web dashboard provides interactive scenario exploration."
    )

    # 5. References
    add_heading("5. References")
    refs = [
        "[1] T. Chen and C. Guestrin, 'XGBoost: A scalable tree boosting system,' in Proc. 22nd ACM SIGKDD Int. Conf. Knowl. Discovery Data Mining, 2016, pp. 785–794.",
        "[2] F. T. Liu, K. M. Ting, and Z.-H. Zhou, 'Isolation forest,' in Proc. 8th IEEE Int. Conf. Data Mining (ICDM), 2008, pp. 413–422.",
        "[3] F. X. Diebold and R. S. Mariano, 'Comparing predictive accuracy,' Journal of Business & Economic Statistics, vol. 13, no. 3, pp. 253–263, 1995."
    ]
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_after = Pt(1.5)
        r_ref = p_ref.add_run(ref)
        r_ref.font.name = 'Book Antiqua'
        r_ref.font.size = Pt(8.0)

    return doc

def generate_abstract_html(anonymized=False):
    author_block = """
        <div class="author-block">
            <strong>Anonymous Author(s) for Double-Blind Peer Review</strong><br>
            <span style="font-size: 8pt; color: #64748b; font-style: italic;">
                (Author names, institutional affiliations, and repository links withheld in compliance with double-blind review policy)
            </span>
        </div>
    """ if anonymized else """
        <div class="author-block">
            <strong>Md Mohi Uddin Mahedi¹, Mst. Mimuna Parvin Mim²</strong><br>
            <span style="font-size: 8pt; color: #475569;">
                Department of Computer Science and Engineering<br>
                Open-Source Research Repository: <u>https://github.com/Mahedi8x/ai_economic</u>
            </span>
        </div>
    """

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Extended Abstract - IEEE CS BDC Symposium 2026</title>
    <style>
        @page {{
            size: A4;
            margin: 25.4mm 19.05mm 25.4mm 19.05mm; /* 1.0" top/bottom, 0.75" left/right */
        }}

        body {{
            font-family: 'Book Antiqua', Palatino, 'Palatino Linotype', Georgia, serif;
            color: #0f172a;
            line-height: 1.34;
            font-size: 9.2pt;
            margin: 0;
            padding: 0;
            background: #ffffff;
            text-align: justify;
        }}

        .header {{
            text-align: center;
            margin-bottom: 8px;
        }}

        h1.title {{
            font-size: 12.5pt;
            font-weight: bold;
            color: #0f172a;
            margin: 0 0 3px 0;
            line-height: 1.22;
            text-align: center;
        }}

        .conf-badge {{
            font-size: 8.5pt;
            font-style: italic;
            color: #475569;
            margin-bottom: 4px;
            text-transform: uppercase;
            letter-spacing: 0.3px;
        }}

        .author-block {{
            text-align: center;
            margin: 4px 0 6px 0;
            font-size: 8.5pt;
        }}

        .keywords {{
            margin-bottom: 8px;
            font-size: 8.8pt;
        }}

        .keywords strong {{
            font-weight: bold;
            color: #1e293b;
        }}

        h2.sec-heading {{
            font-size: 9.8pt;
            font-weight: bold;
            color: #1e3a8a;
            margin-top: 7px;
            margin-bottom: 2px;
            page-break-after: avoid;
        }}

        p {{
            margin: 0 0 5px 0;
            text-indent: 0;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 6px 0 2px 0;
            font-size: 7.8pt;
            page-break-inside: avoid;
        }}

        th {{
            background-color: #1e3a8a;
            color: #ffffff;
            font-weight: bold;
            padding: 3.5px 5px;
            border: 1px solid #94a3b8;
            text-align: center;
        }}

        td {{
            padding: 3px 5px;
            border: 1px solid #cbd5e1;
            text-align: center;
        }}

        tr:nth-child(even) {{
            background-color: #f8fafc;
        }}

        .table-caption {{
            font-size: 7.2pt;
            font-style: italic;
            color: #475569;
            text-align: center;
            margin-bottom: 5px;
        }}

        .ref-item {{
            margin-bottom: 2px;
            font-size: 7.8pt;
            line-height: 1.25;
        }}
    </style>
</head>
<body>

    <div class="header">
        <h1 class="title">Explainable Machine Learning and Grounded Policy Retrieval Framework for Food Inflation Forecasting and Economic Early Warning in Bangladesh</h1>
        <div class="conf-badge">Extended Abstract — IEEE CS BDC Symposium 2026</div>
        {author_block}
    </div>

    <div class="keywords">
        <strong>Keywords—</strong> <em>Food Inflation Nowcasting, Explainable Machine Learning, Economic Early Warning, Conformal Prediction, Policy Information Retrieval.</em>
    </div>

    <h2 class="sec-heading">1. Research Area</h2>
    <p>
        Applied Machine Learning, Computational Macroeconomics, Empirical Time-Series Nowcasting, and Information Retrieval.
    </p>

    <h2 class="sec-heading">2. Goals/Objectives</h2>
    <p>
        This research investigates macroeconomic volatility and staple food price forecasting in developing, import-dependent economies, focusing empirically on Bangladesh. In official surveillance, Consumer Price Index (CPI) and foreign exchange reserve releases from the Bangladesh Bureau of Statistics (BBS) and Bangladesh Bank typically appear around Day 7 to Day 10 of following month (Month M+1). While this represents a modest administrative delay of one to two weeks, rapid international commodity and currency fluctuations during shock regimes create an informational gap for monetary and procurement authorities. The core objective is to evaluate whether non-linear machine learning architectures can extract actionable early signals from high-frequency international markets and domestic macro indicators ahead of official monthly bulletins. Specifically, the study investigates: (i) whether gradient boosting improves short-horizon food inflation nowcasting relative to linear regression and naive persistence baselines; (ii) how empirical conformal prediction behaves when calibrating retail food price uncertainty bands; (iii) whether unsupervised anomaly detection can provide advance warning before systemic macroeconomic distress episodes; and (iv) how policy retrieval pipelines perform when grounding domain economic inquiries on central bank publications.
    </p>

    <h2 class="sec-heading">3. Methodology</h2>
    <p>
        The empirical architecture is anchored on a continuous daily macroeconomic panel spanning January 2010 to September 2026 (6,117 daily observations, 41 indicators) combined with an official monthly macroeconomic panel (201 continuous months, zero synthetic interpolation). Institutional release schedules (BBS CPI on Day 10 of M+1, Bangladesh Bank reserves on Day 7 of M+1, WFP retail food price surveys on Day 15 of M) are strictly enforced via step-function forward fills to prevent look-ahead bias, while daily international shocks (Brent crude oil, CBOT wheat futures, US Dollar Index, natural gas, gold) enter contemporaneously. Annual structural indicators from the World Bank (external debt, import coverage ratios) are incorporated on an as-of release basis, though annual step-filling introduces an open look-ahead timing consideration. For food inflation nowcasting, an Extreme Gradient Boosting (XGBoost) model is trained using chronological walk-forward splits and compared against linear Ridge regression, Random Forest, and naive persistence baselines using Harvey-Leybourne-Newbold corrected Diebold-Mariano tests. Feature attributions are extracted via TreeSHAP. For staple retail commodities (coarse rice, lentils, wheat flour), a rolling walk-forward adaptive conformal prediction engine is evaluated over a 730-day calibration window to generate distribution-free 80% nominal coverage intervals. For crisis surveillance, an Isolation Forest is trained exclusively on pre-2022 historical controls and evaluated out-of-sample on 2022–2026 stress regimes with a reset-capable Cumulative Sum (CUSUM) filter. A prototype dense retrieval module (ChromaDB) indexes authentic central bank policy texts.
    </p>

    <h2 class="sec-heading">4. Results and Analysis</h2>
    <p>
        In monthly out-of-sample evaluation (2024–2026, n &asymp; 33 months), XGBoost reduces root mean square error (RMSE) relative to a linear Ridge baseline by 20.0% at the t+1 horizon (1.171% vs. 1.464%) and by 32.5% at the t+3 horizon (1.740% vs. 2.578%), but neither difference is statistically significant at the 5% level (HLN-corrected Diebold-Mariano p = 0.068, 0.058), and XGBoost does not outperform naive persistence (1.160% at t+1) or Random Forest (1.087% at t+1, 1.725% at t+3) as detailed in Table I. In high-frequency daily 30-day trajectory forecasting, XGBoost substantially outperforms Ridge (RMSE 1.158% vs. 2.457%, a 52.9% reduction; DM = 3.74, p &lt; 0.001), yet remains marginally behind naive persistence (1.137%), indicating dominant short-term price stickiness in domestic staple markets. TreeSHAP feature attribution assigns the highest importance to reserve-adequacy indicators (months of imports: 19.3%), remittance inflows (15.9%), and headline inflation persistence (15.4%), while global crude oil (1.4%) and CBOT wheat (1.2%) exhibit minor direct attribution; annual structural indicators among them remain an open look-ahead concern.
    </p>

    <table>
        <thead>
            <tr>
                <th>Forecast Horizon</th>
                <th>Naive Baseline</th>
                <th>Linear Ridge</th>
                <th>Random Forest</th>
                <th>Proposed XGBoost</th>
                <th>DM Stat vs Ridge (p-val)</th>
            </tr>
        </thead>
        <tbody>
            <tr>
                <td>t+1 (1 Month Ahead)</td>
                <td>1.160%</td>
                <td>1.464%</td>
                <td>1.087%</td>
                <td>1.171% (-20.0%)</td>
                <td>1.82 (p = 0.068)</td>
            </tr>
            <tr>
                <td>t+3 (3 Months Ahead)</td>
                <td>1.757%</td>
                <td>2.578%</td>
                <td>1.725%</td>
                <td>1.740% (-32.5%)</td>
                <td>1.90 (p = 0.058)</td>
            </tr>
            <tr>
                <td>30-Day Forward Daily</td>
                <td>1.137%</td>
                <td>2.457%</td>
                <td>—</td>
                <td><strong>1.158% (-52.9%)</strong></td>
                <td><strong>3.74 (p &lt; 0.001)</strong></td>
            </tr>
        </tbody>
    </table>
    <div class="table-caption">
        Table I: Out-of-Sample Forecasting Performance & Diebold-Mariano Tests across Horizons
    </div>

    <p>
        In retail price uncertainty estimation, rolling walk-forward conformal prediction achieved 85.1% empirical out-of-sample coverage for coarse rice over 1,366 unseen test days, whereas lentils (74.8%) and flour (75.1%) fell short of the 80% nominal target, highlighting commodity-specific calibration disparities. The unsupervised anomaly detector flagged both out-of-sample economic crisis episodes 80 days (2022 foreign exchange drain) and 90 days (2024 political shock) ahead of onset, but flagged 74.0% of out-of-sample days due to persistent post-2022 macroeconomic disequilibrium, indicating that raw alert frequencies must be coupled with secondary persistence filters before operational deployment. The retrieval module was audited for keyword and semantic coverage, achieving an 82.0% mean grounding score across 10 central bank queries; full LLM generation and hallucination benchmarks have not yet been evaluated. A prototype web dashboard provides interactive scenario exploration.
    </p>

    <h2 class="sec-heading">5. References</h2>
    <div class="ref-item">
        [1] T. Chen and C. Guestrin, "XGBoost: A scalable tree boosting system," in <em>Proc. 22nd ACM SIGKDD Int. Conf. Knowl. Discovery Data Mining</em>, 2016, pp. 785–794.
    </div>
    <div class="ref-item">
        [2] F. T. Liu, K. M. Ting, and Z.-H. Zhou, "Isolation forest," in <em>Proc. 8th IEEE Int. Conf. Data Mining (ICDM)</em>, 2008, pp. 413–422.
    </div>
    <div class="ref-item">
        [3] F. X. Diebold and R. S. Mariano, "Comparing predictive accuracy," <em>Journal of Business & Economic Statistics</em>, vol. 13, no. 3, pp. 253–263, 1995.
    </div>

</body>
</html>
"""

def main():
    print("[1/4] Generating DOCX files...")
    doc_anon = build_docx(anonymized=True)
    doc_auth = build_docx(anonymized=False)
    
    anon_docx = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract_Anonymized.docx"
    auth_docx = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract.docx"
    
    doc_anon.save(anon_docx)
    doc_auth.save(auth_docx)
    print(f" [+] Saved: {anon_docx}")
    print(f" [+] Saved: {auth_docx}")

    print("[2/4] Rendering publication-grade PDFs via Edge Headless...")
    # Anonymized HTML & PDF
    anon_html = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract_Anonymized.html"
    anon_pdf = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract_Anonymized.pdf"
    with open(anon_html, "w", encoding="utf-8") as f:
        f.write(generate_abstract_html(anonymized=True))
    
    subprocess.run([
        EDGE_PATH, "--headless", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={anon_pdf}", str(anon_html)
    ], capture_output=True, text=True)

    # Author HTML & PDF
    auth_html = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract.html"
    auth_pdf = BASE_DIR / "IEEE_CS_BDC_Symposium_2026_Extended_Abstract.pdf"
    with open(auth_html, "w", encoding="utf-8") as f:
        f.write(generate_abstract_html(anonymized=False))

    subprocess.run([
        EDGE_PATH, "--headless", "--disable-gpu", "--no-pdf-header-footer",
        f"--print-to-pdf={auth_pdf}", str(auth_html)
    ], capture_output=True, text=True)

    print(f" [+] Rendered: {anon_pdf}")
    print(f" [+] Rendered: {auth_pdf}")

    print("[3/4] Copying all files to Desktop...")
    for f in [anon_docx, auth_docx, anon_pdf, auth_pdf]:
        shutil.copy2(f, DESKTOP_DIR / f.name)
        print(f" [+] Copied to Desktop: {DESKTOP_DIR / f.name}")

    print("[4/4] Verifying page count...")
    import pypdf
    for p in [anon_pdf, auth_pdf]:
        reader = pypdf.PdfReader(str(p))
        print(f" [+] {p.name} -> Exact Page Count: {len(reader.pages)}")

if __name__ == "__main__":
    main()
