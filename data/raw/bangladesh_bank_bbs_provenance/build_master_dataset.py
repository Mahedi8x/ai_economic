import os
import re
import openpyxl
import xlrd
import pandas as pd
import numpy as np

# Directory Paths
base_dir = r"C:\Users\Mehedi\.gemini\antigravity\scratch\bangladesh_macro_real_2010_2026"
source_dir = os.path.join(base_dir, "source_files")
pdf_dir = os.path.join(base_dir, "pdf_reports")
nsdp_dir = r"C:\Users\Mehedi\.gemini\antigravity\scratch\nsdp_files"
scratch_dir = r"C:\Users\Mehedi\.gemini\antigravity\scratch"

os.makedirs(base_dir, exist_ok=True)
os.makedirs(source_dir, exist_ok=True)
os.makedirs(pdf_dir, exist_ok=True)

# 1. Initialize complete 202-month calendar (2010-01 to 2026-10)
dates = pd.date_range("2010-01-01", "2026-10-01", freq="MS")
df = pd.DataFrame(index=dates)
df.index.name = "Date"
df["Date_Str"] = df.index.strftime("%Y-%m")
df["Year"] = df.index.year
df["Month"] = df.index.month
df["Month_Name"] = df.index.strftime("%B")

print(f"Initialized {len(df)} monthly rows from {df['Date_Str'].iloc[0]} to {df['Date_Str'].iloc[-1]}")

# -------------------------------------------------------------
# 2. Extract Official Foreign Exchange Reserves (Gross & BPM6)
# Source: Bangladesh Bank NSDP (Official Reserve Assets - ora_bgd.xlsx)
# -------------------------------------------------------------
ora_path = os.path.join(nsdp_dir, "ora_bgd.xlsx")
wb_ora = openpyxl.load_workbook(ora_path, data_only=True)
ws_ora = wb_ora["New Format"]
rows_ora = list(ws_ora.iter_rows(values_only=True))

date_cols = rows_ora[3]
res_gross_vals = rows_ora[8]
res_bpm6_vals = rows_ora[9]

res_gross_dict = {}
res_bpm6_dict = {}

for col_idx in range(1, len(date_cols)):
    d_str = str(date_cols[col_idx]).strip()
    if re.match(r"^\d{4}-\d{2}$", d_str):
        # Gross
        g_val = res_gross_vals[col_idx]
        if g_val is not None and str(g_val).strip() not in ["---", ""]:
            try:
                res_gross_dict[d_str] = round(float(g_val) / 1000.0, 3)
            except:
                pass
        # BPM6
        b_val = res_bpm6_vals[col_idx]
        if b_val is not None and str(b_val).strip() not in ["---", ""]:
            try:
                res_bpm6_dict[d_str] = round(float(b_val) / 1000.0, 3)
            except:
                pass

# October 2026 update from BB official intreserve / statement of affairs
res_gross_dict["2026-10"] = 36.015  # $36.015 Billion USD
res_bpm6_dict["2026-10"] = 31.199   # $31.199 Billion USD

df["Foreign_Reserves_USD_Bn"] = df["Date_Str"].map(res_gross_dict)
df["Foreign_Reserves_BPM6_USD_Bn"] = df["Date_Str"].map(res_bpm6_dict)

print("Reserves mapped. Missing gross:", df["Foreign_Reserves_USD_Bn"].isna().sum())

# -------------------------------------------------------------
# 3. Extract Official USD/BDT Exchange Rate (Period Average & End of Period)
# Source: Bangladesh Bank NSDP (exr_bgd.xlsx)
# -------------------------------------------------------------
exr_path = os.path.join(nsdp_dir, "exr_bgd.xlsx")
wb_exr = openpyxl.load_workbook(exr_path, data_only=True)
ws_exr = wb_exr["Dataset"]
rows_exr = list(ws_exr.iter_rows(values_only=True))

exr_dates = rows_exr[3]
exr_end_vals = rows_exr[4]
exr_avg_vals = rows_exr[5]

exr_avg_dict = {}
exr_end_dict = {}

for col_idx in range(1, len(exr_dates)):
    d_str = str(exr_dates[col_idx]).strip()
    if re.match(r"^\d{4}-\d{2}$", d_str):
        a_val = exr_avg_vals[col_idx]
        e_val = exr_end_vals[col_idx]
        if a_val is not None:
            try:
                exr_avg_dict[d_str] = round(float(a_val), 4)
            except:
                pass
        if e_val is not None:
            try:
                exr_end_dict[d_str] = round(float(e_val), 4)
            except:
                pass

# October 2026 update from BB official exchangerate
exr_avg_dict["2026-10"] = 123.9900
exr_end_dict["2026-10"] = 124.0500

df["USD_BDT_Exchange_Rate"] = df["Date_Str"].map(exr_avg_dict)
df["USD_BDT_Exchange_Rate_EndPeriod"] = df["Date_Str"].map(exr_end_dict)

print("Exchange rate mapped. Missing:", df["USD_BDT_Exchange_Rate"].isna().sum())

# -------------------------------------------------------------
# 4. Extract Remittance (USD Billion)
# Source: BB NSDP (remit_bgd.xlsx for 2015-2026) & BB Table XXI/XX (2010-2014)
# -------------------------------------------------------------
remit_dict = {}

# A. From remit_bgd.xlsx (2015-01 to 2026-09)
remit_path = os.path.join(nsdp_dir, "remit_bgd.xlsx")
wb_remit = openpyxl.load_workbook(remit_path, data_only=True)
ws_remit = wb_remit.active
rows_remit = list(ws_remit.iter_rows(values_only=True))
remit_dates = rows_remit[3]
remit_vals = rows_remit[4]

for col_idx in range(1, len(remit_dates)):
    d_str = str(remit_dates[col_idx]).strip()
    if re.match(r"^\d{4}-\d{2}$", d_str):
        val = remit_vals[col_idx]
        if val is not None:
            try:
                remit_dict[d_str] = round(float(val) / 1000.0, 3)
            except:
                pass

# B. Historical Remittance for 2010-01 to 2014-12 from BB Table XXI/XX
# 2010: Jan to Dec
remit_dict["2010-01"] = 0.952
remit_dict["2010-02"] = 0.828
remit_dict["2010-03"] = 0.956
remit_dict["2010-04"] = 0.922
remit_dict["2010-05"] = 0.903
remit_dict["2010-06"] = 0.892
remit_dict["2010-07"] = 0.857
remit_dict["2010-08"] = 0.964
remit_dict["2010-09"] = 0.838
remit_dict["2010-10"] = 0.924
remit_dict["2010-11"] = 0.999
remit_dict["2010-12"] = 0.969

# 2011: Jan to Dec
remit_dict["2011-01"] = 0.971
remit_dict["2011-02"] = 0.987
remit_dict["2011-03"] = 1.103
remit_dict["2011-04"] = 1.002
remit_dict["2011-05"] = 0.998
remit_dict["2011-06"] = 1.039
remit_dict["2011-07"] = 1.016
remit_dict["2011-08"] = 1.102
remit_dict["2011-09"] = 0.855
remit_dict["2011-10"] = 1.039
remit_dict["2011-11"] = 0.909
remit_dict["2011-12"] = 1.141

# 2012: Jan to Dec
remit_dict["2012-01"] = 1.221
remit_dict["2012-02"] = 1.131
remit_dict["2012-03"] = 1.225
remit_dict["2012-04"] = 1.082
remit_dict["2012-05"] = 1.147
remit_dict["2012-06"] = 1.058
remit_dict["2012-07"] = 1.201
remit_dict["2012-08"] = 1.179
remit_dict["2012-09"] = 1.104
remit_dict["2012-10"] = 1.454
remit_dict["2012-11"] = 1.101
remit_dict["2012-12"] = 1.284

# 2013: Jan to Dec
remit_dict["2013-01"] = 1.327
remit_dict["2013-02"] = 1.160
remit_dict["2013-03"] = 1.230
remit_dict["2013-04"] = 1.191
remit_dict["2013-05"] = 1.087
remit_dict["2013-06"] = 1.144
remit_dict["2013-07"] = 1.239
remit_dict["2013-08"] = 1.006
remit_dict["2013-09"] = 1.026
remit_dict["2013-10"] = 1.227
remit_dict["2013-11"] = 1.059
remit_dict["2013-12"] = 1.214

# 2014: Jan to Dec
remit_dict["2014-01"] = 1.261
remit_dict["2014-02"] = 1.162
remit_dict["2014-03"] = 1.282
remit_dict["2014-04"] = 1.235
remit_dict["2014-05"] = 1.205
remit_dict["2014-06"] = 1.313
remit_dict["2014-07"] = 1.492
remit_dict["2014-08"] = 1.174
remit_dict["2014-09"] = 1.341
remit_dict["2014-10"] = 1.018
remit_dict["2014-11"] = 1.183
remit_dict["2014-12"] = 1.275

# October 2026 update
remit_dict["2026-10"] = 2.562

df["Remittance_USD_Bn"] = df["Date_Str"].map(remit_dict)
print("Remittance mapped. Missing:", df["Remittance_USD_Bn"].isna().sum())

# -------------------------------------------------------------
# 5. Extract Headline Inflation (%) & Food Inflation (%)
# Source: BBS / Bangladesh Bank Table VII & Table VIII
# -------------------------------------------------------------
headline_dict = {}
food_dict = {}
cpi_base_dict = {}

# A. 2010-01 to 2012-06 (Base 1995-96=100) from stat_dec11.xls and stat_dec14.xls
inf_pre2012 = {
    # 2010
    "2010-01": (8.99, 10.56, "1995-96=100"),
    "2010-02": (9.06, 10.93, "1995-96=100"),
    "2010-03": (8.78, 10.80, "1995-96=100"),
    "2010-04": (8.54, 10.47, "1995-96=100"),
    "2010-05": (8.65, 10.72, "1995-96=100"),
    "2010-06": (8.70, 10.88, "1995-96=100"),
    "2010-07": (7.26, 8.72, "1995-96=100"),
    "2010-08": (7.52, 9.64, "1995-96=100"),
    "2010-09": (7.61, 9.72, "1995-96=100"),
    "2010-10": (6.86, 8.43, "1995-96=100"),
    "2010-11": (7.54, 9.80, "1995-96=100"),
    "2010-12": (8.28, 11.01, "1995-96=100"),
    # 2011
    "2011-01": (9.04, 11.91, "1995-96=100"),
    "2011-02": (9.79, 12.77, "1995-96=100"),
    "2011-03": (10.49, 13.87, "1995-96=100"),
    "2011-04": (10.67, 14.36, "1995-96=100"),
    "2011-05": (10.20, 13.16, "1995-96=100"),
    "2011-06": (10.17, 12.51, "1995-96=100"),
    "2011-07": (10.96, 13.40, "1995-96=100"),
    "2011-08": (11.29, 12.70, "1995-96=100"),
    "2011-09": (11.97, 13.75, "1995-96=100"),
    "2011-10": (11.42, 12.82, "1995-96=100"),
    "2011-11": (11.58, 12.47, "1995-96=100"),
    "2011-12": (10.63, 10.40, "1995-96=100"),
    # 2012 (Jan - Jun)
    "2012-01": (11.59, 11.83, "1995-96=100"),
    "2012-02": (10.43, 8.92, "1995-96=100"),
    "2012-03": (10.10, 8.28, "1995-96=100"),
    "2012-04": (9.93, 7.82, "1995-96=100"),
    "2012-05": (9.15, 6.45, "1995-96=100"),
    "2012-06": (8.56, 5.71, "1995-96=100"),
}

for k, (h, f, b) in inf_pre2012.items():
    headline_dict[k] = h
    food_dict[k] = f
    cpi_base_dict[k] = b

# B. 2012-07 to 2023-03 (Base 2005-06=100) from BBS / BB Table VII
inf_2005_base = {
    # 2012 (Jul - Dec)
    "2012-07": (5.21, 2.23), "2012-08": (4.97, 2.25), "2012-09": (4.96, 1.75),
    "2012-10": (5.86, 2.51), "2012-11": (6.55, 3.94), "2012-12": (7.14, 5.28),
    # 2013
    "2013-01": (6.62, 5.02), "2013-02": (7.84, 7.45), "2013-03": (7.71, 7.50),
    "2013-04": (8.37, 8.68), "2013-05": (7.98, 8.13), "2013-06": (8.05, 8.26),
    "2013-07": (7.85, 8.14), "2013-08": (7.39, 8.09), "2013-09": (7.13, 7.93),
    "2013-10": (7.03, 8.38), "2013-11": (7.15, 8.55), "2013-12": (7.35, 9.00),
    # 2014
    "2014-01": (7.50, 8.81), "2014-02": (7.44, 8.84), "2014-03": (7.48, 8.96),
    "2014-04": (7.46, 8.95), "2014-05": (7.48, 9.09), "2014-06": (6.97, 8.00),
    "2014-07": (7.04, 7.94), "2014-08": (6.91, 7.67), "2014-09": (6.84, 7.63),
    "2014-10": (6.60, 7.16), "2014-11": (6.21, 6.44), "2014-12": (6.11, 5.86),
    # 2015
    "2015-01": (6.04, 6.07), "2015-02": (6.14, 6.11), "2015-03": (6.27, 6.37),
    "2015-04": (6.32, 6.48), "2015-05": (6.19, 6.23), "2015-06": (6.25, 6.32),
    "2015-07": (6.36, 6.07), "2015-08": (6.17, 6.06), "2015-09": (6.24, 5.92),
    "2015-10": (6.19, 5.89), "2015-11": (6.05, 5.72), "2015-12": (6.10, 5.48),
    # 2016
    "2016-01": (6.07, 4.33), "2016-02": (5.62, 3.77), "2016-03": (5.65, 3.89),
    "2016-04": (5.61, 3.84), "2016-05": (5.45, 3.81), "2016-06": (5.53, 4.23),
    "2016-07": (5.40, 4.35), "2016-08": (5.37, 4.30), "2016-09": (5.53, 5.10),
    "2016-10": (5.57, 5.56), "2016-11": (5.38, 5.41), "2016-12": (5.03, 5.38),
    # 2017
    "2017-01": (5.15, 6.53), "2017-02": (5.31, 6.84), "2017-03": (5.39, 6.89),
    "2017-04": (5.47, 6.94), "2017-05": (5.44, 6.90), "2017-06": (5.94, 7.51),
    "2017-07": (5.57, 6.95), "2017-08": (6.02, 7.28), "2017-09": (6.12, 7.87),
    "2017-10": (6.04, 7.62), "2017-11": (5.91, 7.09), "2017-12": (5.83, 7.13),
    # 2018
    "2018-01": (5.88, 7.62), "2018-02": (5.72, 7.27), "2018-03": (5.68, 7.09),
    "2018-04": (5.63, 7.03), "2018-05": (5.57, 6.56), "2018-06": (5.54, 5.98),
    "2018-07": (5.51, 6.18), "2018-08": (5.48, 5.97), "2018-09": (5.43, 5.42),
    "2018-10": (5.40, 5.14), "2018-11": (5.37, 5.29), "2018-12": (5.35, 5.28),
    # 2019
    "2019-01": (5.42, 5.33), "2019-02": (5.47, 5.44), "2019-03": (5.55, 5.63),
    "2019-04": (5.58, 5.61), "2019-05": (5.63, 5.66), "2019-06": (5.52, 5.40),
    "2019-07": (5.62, 5.42), "2019-08": (5.49, 5.27), "2019-09": (5.54, 5.30),
    "2019-10": (5.47, 5.49), "2019-11": (6.05, 6.41), "2019-12": (5.75, 5.88),
    # 2020
    "2020-01": (5.57, 5.12), "2020-02": (5.46, 4.97), "2020-03": (5.48, 4.87),
    "2020-04": (5.96, 5.91), "2020-05": (5.35, 5.09), "2020-06": (6.02, 6.54),
    "2020-07": (5.53, 5.70), "2020-08": (5.68, 6.08), "2020-09": (5.97, 6.50),
    "2020-10": (6.44, 7.34), "2020-11": (5.52, 5.73), "2020-12": (5.29, 5.34),
    # 2021
    "2021-01": (5.02, 5.23), "2021-02": (5.32, 5.42), "2021-03": (5.47, 5.51),
    "2021-04": (5.56, 5.57), "2021-05": (5.26, 4.87), "2021-06": (5.64, 5.45),
    "2021-07": (5.36, 5.08), "2021-08": (5.54, 5.16), "2021-09": (5.59, 5.21),
    "2021-10": (5.70, 5.22), "2021-11": (5.98, 5.43), "2021-12": (6.05, 5.46),
    # 2022
    "2022-01": (5.86, 5.60), "2022-02": (6.17, 6.10), "2022-03": (6.22, 6.34),
    "2022-04": (6.29, 6.24), "2022-05": (7.42, 8.30), "2022-06": (7.56, 8.37),
    "2022-07": (7.48, 8.19), "2022-08": (9.52, 9.94), "2022-09": (9.10, 9.08),
    "2022-10": (8.91, 8.50), "2022-11": (8.85, 8.14), "2022-12": (8.71, 7.91),
    # 2023 (Jan - Mar)
    "2023-01": (8.57, 7.76), "2023-02": (8.78, 8.13), "2023-03": (9.33, 9.09)
}

for k, (h, f) in inf_2005_base.items():
    headline_dict[k] = h
    food_dict[k] = f
    cpi_base_dict[k] = "2005-06=100"

# C. 2023-04 to 2026-10 (Base 2021-22=100) from BBS / BB Table VIII
inf_2021_base = {
    # 2023 (Apr - Dec)
    "2023-04": (9.24, 8.84), "2023-05": (9.94, 9.24), "2023-06": (9.74, 9.73),
    "2023-07": (9.69, 9.76), "2023-08": (9.92, 12.54), "2023-09": (9.63, 12.37),
    "2023-10": (9.93, 12.56), "2023-11": (9.49, 10.76), "2023-12": (9.41, 9.58),
    # 2024
    "2024-01": (9.86, 9.56), "2024-02": (9.67, 9.44), "2024-03": (9.81, 9.87),
    "2024-04": (9.74, 10.22), "2024-05": (9.89, 10.79), "2024-06": (9.72, 10.42),
    "2024-07": (11.66, 14.10), "2024-08": (10.49, 11.36), "2024-09": (9.92, 10.40),
    "2024-10": (10.87, 12.66), "2024-11": (11.38, 13.80), "2024-12": (10.89, 12.92),
    # 2025
    "2025-01": (9.94, 10.72), "2025-02": (9.32, 9.24), "2025-03": (9.35, 8.93),
    "2025-04": (9.17, 8.63), "2025-05": (9.05, 8.59), "2025-06": (8.48, 7.39),
    "2025-07": (8.55, 7.56), "2025-08": (8.29, 7.60), "2025-09": (8.36, 7.64),
    "2025-10": (8.17, 7.08), "2025-11": (8.29, 7.36), "2025-12": (8.49, 7.71),
    # 2026
    "2026-01": (8.58, 8.29), "2026-02": (9.13, 9.30), "2026-03": (8.71, 8.24),
    "2026-04": (9.04, 8.39), "2026-05": (9.42, 9.06), "2026-06": (9.16, 8.60),
    "2026-07": (8.32, 7.16), "2026-08": (8.26, 7.02), "2026-09": (8.34, 7.22),
    "2026-10": (8.28, 7.15)
}

for k, (h, f) in inf_2021_base.items():
    headline_dict[k] = h
    food_dict[k] = f
    cpi_base_dict[k] = "2021-22=100"

df["Headline_Inflation_pct"] = df["Date_Str"].map(headline_dict)
df["Food_Inflation_pct"] = df["Date_Str"].map(food_dict)
df["CPI_Base_Year"] = df["Date_Str"].map(cpi_base_dict)

print("Headline Inflation mapped. Missing:", df["Headline_Inflation_pct"].isna().sum())
print("Food Inflation mapped. Missing:", df["Food_Inflation_pct"].isna().sum())

# -------------------------------------------------------------
# 6. Add Official Verification URLs & Source Document
# -------------------------------------------------------------
def assign_source_info(row):
    year = row["Year"]
    d_str = row["Date_Str"]
    if year == 2010:
        pub = "BB Monthly Economic Trends (Oct/Dec 2010)"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/oct10/statisticaltable.pdf"
    elif year == 2011:
        pub = "BB Monthly Economic Trends (Dec 2011)"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec11/statisticaltable.pdf"
    elif year == 2012:
        pub = "BB Monthly Economic Trends (Dec 2012)"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec12/econtrds.php"
    elif year == 2013:
        pub = "BB Monthly Economic Trends (Dec 2013)"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec13/statisticaltable.pdf"
    elif year == 2014:
        pub = "BB Monthly Economic Trends (Dec 2014)"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec14/statisticaltable.pdf"
    elif year == 2015:
        pub = "BB Monthly Economic Trends (Dec 2015) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec15/statisticaltable.pdf"
    elif year == 2016:
        pub = "BB Monthly Economic Trends (Dec 2016) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec16/statisticaltable.pdf"
    elif year == 2017:
        pub = "BB Monthly Economic Trends (Dec 2017) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec17/statisticaltable.pdf"
    elif year == 2018:
        pub = "BB Monthly Economic Trends (Dec 2018) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec18/econtrds.php"
    elif year == 2019:
        pub = "BB Monthly Economic Trends (Dec 2019) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec19/statisticaltable.pdf"
    elif year == 2020:
        pub = "BB Monthly Economic Trends (Dec 2020) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec20/statisticaltable.pdf"
    elif year == 2021:
        pub = "BB Monthly Economic Trends (Dec 2021) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/dec21/statisticaltable.pdf"
    elif year == 2022:
        pub = "BB Monthly Economic Trends (Dec 2022) & NSDP"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/etdec22.pdf"
    elif year == 2023:
        pub = "BB Monthly Economic Trends (Dec 2023) & BBS CPI"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/etdec23.pdf"
    elif year == 2024:
        pub = "BB Monthly Economic Trends (Dec 2024) & BBS CPI"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/etdec24.pdf"
    elif year == 2025:
        pub = "BB Monthly Economic Trends (Dec 2025) & BBS CPI"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/etdec25.pdf"
    else: # 2026
        pub = "BB Monthly Economic Trends (Sep 2026) / BB MEI / BBS"
        url = "https://www.bb.org.bd/pub/monthly/econtrds/etsep26.pdf"
    return pd.Series([pub, url])

df[["Source_Document", "Source_URL"]] = df.apply(assign_source_info, axis=1)

# Format Final Columns
columns_order = [
    "Date_Str",
    "Year",
    "Month",
    "Month_Name",
    "Food_Inflation_pct",
    "Headline_Inflation_pct",
    "Foreign_Reserves_USD_Bn",
    "Foreign_Reserves_BPM6_USD_Bn",
    "Remittance_USD_Bn",
    "USD_BDT_Exchange_Rate",
    "USD_BDT_Exchange_Rate_EndPeriod",
    "CPI_Base_Year",
    "Source_Document",
    "Source_URL"
]

df_final = df[columns_order].rename(columns={"Date_Str": "Date"})

# Save clean CSV in target folder
out_csv_path = os.path.join(base_dir, "bangladesh_bank_bbs_monthly_macro_2010_2026_real.csv")
df_final.to_csv(out_csv_path, index=False)
print(f"Master Clean Dataset Saved: {out_csv_path} (Shape: {df_final.shape})")

# Also copy raw/clean version to bangladesh_economic_intelligence folder for compatibility
raw_compat_path = r"C:\Users\Mehedi\.gemini\antigravity\scratch\bangladesh_economic_intelligence\data\bangladesh_bank_bbs_monthly_macro_2010_2026_raw.csv"
df_final.to_csv(raw_compat_path, index=False)
print(f"Compatibility CSV Saved: {raw_compat_path}")

# Verification summary
print("\n=== DATA INTEGRITY & AUDIT REPORT ===")
print(f"Total Rows: {len(df_final)}")
print(f"Date Range: {df_final['Date'].iloc[0]} to {df_final['Date'].iloc[-1]}")
print(f"Null count in Food_Inflation_pct: {df_final['Food_Inflation_pct'].isna().sum()}")
print(f"Null count in Headline_Inflation_pct: {df_final['Headline_Inflation_pct'].isna().sum()}")
print(f"Null count in Foreign_Reserves_USD_Bn: {df_final['Foreign_Reserves_USD_Bn'].isna().sum()}")
print(f"Null count in Remittance_USD_Bn: {df_final['Remittance_USD_Bn'].isna().sum()}")
print(f"Null count in USD_BDT_Exchange_Rate: {df_final['USD_BDT_Exchange_Rate'].isna().sum()}")
print("\nSample Data (First 5 Months - 2010):")
print(df_final[["Date", "Food_Inflation_pct", "Headline_Inflation_pct", "Foreign_Reserves_USD_Bn", "Remittance_USD_Bn", "USD_BDT_Exchange_Rate"]].head())
print("\nSample Data (Last 5 Months - 2026):")
print(df_final[["Date", "Food_Inflation_pct", "Headline_Inflation_pct", "Foreign_Reserves_USD_Bn", "Remittance_USD_Bn", "USD_BDT_Exchange_Rate"]].tail())
